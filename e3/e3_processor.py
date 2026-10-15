"""
e3_processor.py
===============
Επεξεργασία αρχείου Excel Ισοζυγίου (Γ Κατηγορίας) για εξαγωγή ποσών Ε3.

Υποστηρίζει ανοικτά και κλειστά (χρεοπιστωμένα) ισοζύγια:
  - Ανοικτό  : ποσό Ε3 = ABS(Υπόλοιπο)
  - Κλειστό  : ομάδα-6 → Χρέωση, ομάδα-7 → Πίστωση

Χρήση:
    from e3_processor import process_excel_file
    result = process_excel_file("/tmp/isozygio.xls")
    # result = {"ok": True, "is_closed_balance_sheet": ..., "e3_totals": [...], ...}
"""

import os
import pandas as pd
import numpy as np

from .e3_field_map import lookup as _e3_lookup, group_by_table as _e3_group


# ─────────────────────────────────────────────────────────────────────────────
# Εσωτερικές βοηθητικές συναρτήσεις
# ─────────────────────────────────────────────────────────────────────────────

def _load_dataframe(excel_path: str):
    """
    Φορτώνει το Excel και εντοπίζει αυτόματα τη γραμμή επικεφαλίδων
    αναζητώντας τις λέξεις-κλειδιά «κωδικός» και «υπόλοιπο».

    Επιστρέφει DataFrame ή None αν δεν βρεθεί έγκυρη γραμμή.
    """
    # with: το αρχείο κλείνει (στα Windows ένα ανοιχτό ExcelFile μπλόκαρε τη διαγραφή του προσωρινού αρχείου)
    with pd.ExcelFile(excel_path) as xls:
        for sheet in xls.sheet_names:
            raw = pd.read_excel(xls, sheet_name=sheet, header=None)
            for i, row_vals in raw.iterrows():
                row_str = " ".join(str(v) for v in row_vals if pd.notna(v)).lower()
                if "κωδικός" in row_str and ("υπόλοιπο" in row_str or "υπολοιπο" in row_str):
                    return pd.read_excel(xls, sheet_name=sheet, header=i)
    return None


def _normalise_columns(df: pd.DataFrame):
    """
    Αντιστοιχίζει τις ελληνικές επικεφαλίδες σε εσωτερικά ονόματα
    (code, desc, debit, credit, balance, e3).

    Επιστρέφει (DataFrame με νέα ονόματα, λίστα λείπουσων στηλών).
    """
    col_map = {}
    for c in df.columns:
        cs = str(c).strip().lower()
        if "κωδικός" in cs or "κωδικος" in cs:
            col_map["code"] = c
        elif "περιγραφή" in cs or "περιγραφη" in cs:
            col_map["desc"] = c
        elif "χρέωση" in cs or "χρεωση" in cs:
            col_map["debit"] = c
        elif "πίστωση" in cs or "πιστωση" in cs:
            col_map["credit"] = c
        elif "υπόλοιπο" in cs or "υπολοιπο" in cs:
            col_map["balance"] = c
        elif "ε3" in cs or "εντύπου" in cs or "εντυπου" in cs:
            col_map["e3"] = c

    required = ["code", "debit", "credit", "balance", "e3"]
    missing = [k for k in required if k not in col_map]

    rename = {v: k for k, v in col_map.items()}
    df = df.rename(columns=rename)

    # Καθαρισμός & μετατροπή τύπων
    df["code"] = df["code"].astype(str).str.strip()
    for num_col in ("debit", "credit", "balance"):
        df[num_col] = pd.to_numeric(df[num_col], errors="coerce").fillna(0.0)
    df["e3"] = pd.to_numeric(df["e3"], errors="coerce")

    # Απόρριψη γραμμών χωρίς κωδικό
    df = df[df["code"].str.len() > 0]
    df = df[df["code"].str.lower() != "nan"]

    return df, missing


def _detect_closed_balance(df: pd.DataFrame) -> bool:
    """
    Ανιχνεύει αν το ισοζύγιο είναι κλειστό (χρεοπιστωμένο).

    Κριτήριο: ≥50% των leaf-λογαριασμών ομάδας-6 έχουν:
      Υπόλοιπο ≈ 0  ΚΑΙ  Χρέωση > 0  ΚΑΙ  Πίστωση > 0
    """
    grp6_leaf = df[
        df["code"].str.startswith("6") &
        df["code"].str.contains("-")
    ]
    if len(grp6_leaf) == 0:
        return False

    closed_mask = (
        (grp6_leaf["balance"].abs() < 0.01) &
        (grp6_leaf["debit"] > 0) &
        (grp6_leaf["credit"] > 0)
    )
    return closed_mask.sum() / len(grp6_leaf) >= 0.5


def _extract_e3_amount(row: pd.Series, is_closed: bool):
    """
    Υπολογίζει το ποσό Ε3 για μία leaf-γραμμή.

    Επιστρέφει dict {e3_code, amount} ή None αν η γραμμή δεν έχει κωδικό Ε3.
    """
    e3_val = row["e3"]
    if pd.isna(e3_val) or e3_val == -1:
        return None

    first_char = row["code"][0] if row["code"] else ""
    try:
        account_group = int(first_char)
    except ValueError:
        account_group = 0

    if is_closed:
        if account_group == 6:
            amount = abs(row["debit"])
        elif account_group == 7:
            amount = abs(row["credit"])
        else:
            amount = abs(row["balance"])
    else:
        amount = abs(row["balance"])

    return {"e3_code": float(e3_val), "amount": amount}


def _enrich_totals(e3_totals: dict) -> tuple:
    """
    Εμπλουτίζει τα σύνολα Ε3 με περιγραφές από το e3_field_map.

    Επιστρέφει (totals_list, grouped_by_table).
    """
    totals_list = []
    for k, v in sorted(e3_totals.items()):
        entry = {"e3_code": k, "amount": round(v, 2)}
        info = _e3_lookup(k)
        if info:
            entry["label"]       = info["label"]
            entry["sub_label"]   = info["sub_label"]
            entry["table"]       = info["table"]
            entry["activity"]    = info["activity"]
            entry["description"] = info["full_description"]
        totals_list.append(entry)

    grouped = _e3_group(totals_list)
    return totals_list, grouped


# ─────────────────────────────────────────────────────────────────────────────
# Δημόσια API
# ─────────────────────────────────────────────────────────────────────────────

def extract_account_leaf_rows(excel_path: str) -> dict:
    """
    Επεξεργάζεται ένα αρχείο Excel ισοζυγίου (ίδια μορφή με process_excel_file)
    αλλά επιστρέφει ΚΑΘΕ leaf-λογαριασμό με μη μηδενικό ποσό, ανεξάρτητα από το αν
    έχει οριζόμενο κωδικό Ε3 — χρησιμοποιείται από το "Λογιστικό Αποτέλεσμα" για να
    αντικαταστήσει (override) τα αυτόματα υπολογισμένα από myDATA ποσά ανά λογαριασμό.

    Returns:
        Dict με ok, is_closed_balance_sheet, leaf_rows: [{account, description, amount, side}]
    """
    df = _load_dataframe(excel_path)
    if df is None:
        return {
            "ok": False,
            "error": (
                "Δεν βρέθηκαν οι στήλες Κωδικός / Υπόλοιπο. "
                "Βεβαιωθείτε ότι το αρχείο είναι ισοζύγιο (Γ κατηγορίας)."
            ),
        }

    df, missing = _normalise_columns(df)
    if missing:
        return {
            "ok": False,
            "error": (
                f"Λείπουν στήλες: {missing}. "
                "Αναμένονται: Κωδικός, Χρέωση, Πίστωση, Υπόλοιπο, Πεδίο Ε3."
            ),
        }

    is_closed = _detect_closed_balance(df)

    leaf_df = df[df["code"].str.contains("-")].copy()
    if len(leaf_df) == 0:
        leaf_df = df.copy()

    rows = []
    for _, row in leaf_df.iterrows():
        first_char = row["code"][0] if row["code"] else ""
        try:
            account_group = int(first_char)
        except ValueError:
            account_group = 0

        if is_closed:
            if account_group == 6:
                amount, side = abs(row["debit"]), "debit"
            elif account_group == 7:
                amount, side = abs(row["credit"]), "credit"
            else:
                amount, side = abs(row["balance"]), "balance"
        else:
            amount, side = abs(row["balance"]), "balance"

        if amount == 0:
            continue

        rows.append({
            "account": row["code"],
            "description": str(row.get("desc", "")).strip() if "desc" in row else "",
            "amount": round(float(amount), 2),
            "side": side,
        })

    return {"ok": True, "is_closed_balance_sheet": is_closed, "leaf_rows": rows}


# Οι στήλες που διαβάζει η εφαρμογή (ίδιες με το _normalise_columns) — για το ⓘ και το πρότυπο Excel.
ISOZYGIO_COLUMNS = [
    {"name": "Κωδικός", "required": True,
     "help": "Κωδικός λογαριασμού. Οι αναλυτικοί λογαριασμοί έχουν παύλες (π.χ. 64-07-00-0000) και αυτοί αθροίζονται· "
             "οι συγκεντρωτικοί (π.χ. 64) αγνοούνται όταν υπάρχουν αναλυτικοί."},
    {"name": "Περιγραφή", "required": False, "help": "Περιγραφή λογαριασμού (προαιρετική, εμφανίζεται στην ανάλυση)."},
    {"name": "Χρέωση", "required": True, "help": "Σύνολο χρεώσεων της χρήσης (αριθμός)."},
    {"name": "Πίστωση", "required": True, "help": "Σύνολο πιστώσεων της χρήσης (αριθμός)."},
    {"name": "Υπόλοιπο", "required": True,
     "help": "Υπόλοιπο λογαριασμού. Ανοικτό ισοζύγιο: ποσό Ε3 = |Υπόλοιπο|. Κλειστό (χρεοπιστωμένο): ομάδα 6 → Χρέωση, ομάδα 7 → Πίστωση."},
    {"name": "Πεδίο Ε3", "required": True,
     "help": "Κωδικός Ε3 του λογαριασμού ως αριθμός με 3 δεκαδικά (π.χ. 585.007, 561.003, 588). "
             "Κενό ή -1 = ο λογαριασμός δεν πάει στο Ε3."},
]

_ISOZYGIO_SAMPLE = [
    ("60", "Αμοιβές και έξοδα προσωπικού", 10000.00, 0.00, 10000.00, None),
    ("60-00-00-0000", "Μικτές αποδοχές", 10000.00, 0.00, 10000.00, 581.001),
    ("64", "Διάφορα έξοδα", 10524.72, 0.00, 10524.72, None),
    ("64-07-00-0000", "Ασφαλιστικές εισφορές επιτηδευματία (e-ΕΦΚΑ/ΤΕΚΑ)", 4064.30, 0.00, 4064.30, 585.007),
    ("64-09-00-0000", "Ενοίκια", 6000.00, 0.00, 6000.00, 585.014),
    ("64-98-00-0000", "Πρόσθετα τέλη και προσαυξήσεις εισφορών", 460.42, 0.00, 460.42, 588),
    ("70", "Πωλήσεις εμπορευμάτων", 0.00, 50000.00, -50000.00, None),
    ("70-00-00-0000", "Πωλήσεις εμπορευμάτων λιανικές", 0.00, 50000.00, -50000.00, 561.003),
    ("38-00-00-0000", "Ταμείο", 60000.00, 40000.00, 20000.00, -1),
]


def build_isozygio_template() -> bytes:
    """Πρότυπο Excel ισοζυγίου (.xlsx) με τις στήλες που διαβάζει η εφαρμογή, παραδείγματα και οδηγίες."""
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Ισοζύγιο"
    ws.append(["Ισοζύγιο Γενικής Λογιστικής — πρότυπο ScanmyData (Έλεγχος Ε3)"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Συμπληρώστε από τη γραμμή 4 και κάτω. Οι επικεφαλίδες της γραμμής 3 πρέπει να μείνουν ως έχουν."])
    ws["A2"].font = Font(italic=True, color="666666")
    ws.append([c["name"] for c in ISOZYGIO_COLUMNS])
    head_fill = PatternFill("solid", fgColor="1E3A8A")
    for cell in ws[3]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in _ISOZYGIO_SAMPLE:
        ws.append(list(row))
    for r in ws.iter_rows(min_row=4, max_row=ws.max_row):
        for c in r[2:5]:
            c.number_format = "#,##0.00"
        r[5].number_format = "0.000"
        if "-" not in str(r[0].value or ""):
            for c in r:
                c.font = Font(bold=True, color="475569")
    for col, w in zip("ABCDEF", (18, 52, 14, 14, 14, 12)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"
    dv = DataValidation(type="decimal", operator="between", formula1="-1", formula2="999.999", allow_blank=True,
                        showErrorMessage=True, errorTitle="Πεδίο Ε3",
                        error="Δώστε κωδικό Ε3 ως αριθμό (π.χ. 585.007) ή -1 / κενό για λογαριασμό εκτός Ε3.")
    ws.add_data_validation(dv)
    dv.add("F4:F5000")

    info = wb.create_sheet("Οδηγίες")
    info.append(["Στήλη", "Υποχρεωτική", "Τι περιέχει"])
    for cell in info[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
    for c in ISOZYGIO_COLUMNS:
        info.append([c["name"], "Ναι" if c["required"] else "Όχι", c["help"]])
    info.append([])
    for line in (
        "Η εφαρμογή βρίσκει μόνη της τη γραμμή επικεφαλίδων (αυτή που περιέχει «Κωδικός» και «Υπόλοιπο»), σε οποιοδήποτε φύλλο.",
        "Δεκτά αρχεία: .xlsx και .xls. Τα ποσά ως αριθμοί (όχι κείμενο).",
        "Οι γραμμές-παραδείγματα στο φύλλο «Ισοζύγιο» μπορούν να σβηστούν.",
        "Συγκρίσεις: 585.007 με τις φορολογικές βεβαιώσεις e-ΕΦΚΑ/ΤΕΚΑ (+ λοιπά ταμεία), 588 με τις επιβαρύνσεις εισφορών, "
        "585.014 με το ενοίκιο/ιδιόχρηση, και όλοι οι κωδικοί με το myDATA.",
    ):
        info.append([line])
    info.column_dimensions["A"].width = 16
    info.column_dimensions["B"].width = 13
    info.column_dimensions["C"].width = 110
    for r in info.iter_rows(min_row=2):
        for c in r:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def process_excel_file(excel_path: str) -> dict:
    """
    Επεξεργάζεται ένα αρχείο Excel ισοζυγίου και εξάγει τα ποσά Ε3.

    Args:
        excel_path: Απόλυτο μονοπάτι προς το .xls/.xlsx αρχείο.

    Returns:
        Dict με κλειδιά:
          ok                    – True αν η επεξεργασία πέτυχε
          error                 – Μήνυμα σφάλματος (μόνο αν ok=False)
          is_closed_balance_sheet – bool
          total_rows            – συνολικές γραμμές μετά φόρτωση
          leaf_rows_processed   – γραμμές που επεξεργάστηκαν (leaf)
          e3_totals             – λίστα {e3_code, amount, label?, table?,...}
          e3_by_table           – dict ομαδοποιημένος ανά πίνακα Ε3
          detail                – ανά-λογαριασμός λεπτομέρεια
    """
    # 1. Φόρτωση
    df = _load_dataframe(excel_path)
    if df is None:
        return {
            "ok": False,
            "error": (
                "Δεν βρέθηκαν οι στήλες Κωδικός / Υπόλοιπο. "
                "Βεβαιωθείτε ότι το αρχείο είναι ισοζύγιο (Γ κατηγορίας)."
            ),
        }

    # 2. Κανονικοποίηση στηλών
    df, missing = _normalise_columns(df)
    if missing:
        return {
            "ok": False,
            "error": (
                f"Λείπουν στήλες: {missing}. "
                "Αναμένονται: Κωδικός, Χρέωση, Πίστωση, Υπόλοιπο, Πεδίο Ε3."
            ),
        }

    # 3. Ανίχνευση κλειστού ισοζυγίου
    is_closed = _detect_closed_balance(df)

    # 4. Επεξεργασία leaf-γραμμών
    leaf_df = df[df["code"].str.contains("-")].copy()
    if len(leaf_df) == 0:
        leaf_df = df.copy()

    e3_totals = {}
    detail_rows = []
    for _, row in leaf_df.iterrows():
        result_item = _extract_e3_amount(row, is_closed)
        if result_item is None:
            continue
        code_key = result_item["e3_code"]
        e3_totals[code_key] = e3_totals.get(code_key, 0.0) + result_item["amount"]
        detail_rows.append({
            "account":     row["code"],
            "description": str(row.get("desc", "")).strip() if "desc" in row else "",
            "e3_code":     code_key,
            "amount":      result_item["amount"],
        })

    # 5. Εμπλουτισμός με περιγραφές & ομαδοποίηση
    totals_list, grouped = _enrich_totals(e3_totals)

    return {
        "ok":                     True,
        "is_closed_balance_sheet": is_closed,
        "total_rows":             len(df),
        "leaf_rows_processed":    len(leaf_df),
        "e3_totals":              totals_list,
        "e3_by_table":            grouped,
        "detail":                 detail_rows,
    }
