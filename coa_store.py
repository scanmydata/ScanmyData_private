# -*- coding: utf-8 -*-
"""Λογιστικό σχέδιο (Β/Γ κατηγορίας) ανά ομάδα ΚΑΙ ανά credential, πρότυπα Excel και έλεγχος αρχείων.

Αποθήκευση
    * κοινό σχέδιο ομάδας (όπως πάντα):  <group>/chart_of_accounts_g.xlsx , chart_of_accounts_b.xlsx
    * ειδικό σχέδιο credential (νέο):    <group>/coa/<ΑΦΜ>_G.xlsx , <ΑΦΜ>_B.xlsx  (+ .meta.json)

Επίλυση (``resolve_coa_path``): πρώτα το ειδικό σχέδιο του credential για την κατηγορία του
(Β ή Γ), αλλιώς το κοινό της ομάδας — άρα όποιος δεν ανεβάσει ειδικό σχέδιο δεν αλλάζει σε τίποτα.

Το module δεν εξαρτάται από Flask/app: όλες οι συναρτήσεις παίρνουν ρητά τον φάκελο της ομάδας.
"""
from __future__ import annotations

import io
import json
import os
import re
import unicodedata
from datetime import datetime
from typing import Any, Dict, Iterable, List, Optional, Tuple

COA_DIRNAME = "coa"
COA_COLUMNS = ("Κωδικός", "Περιγραφή", "Ποσοστό ΦΠΑ", "Λογαριασμός ΦΠΑ")
# Γ: χρειάζονται όλες (ο λογαριασμός ΦΠΑ προστίθεται αυτόματα στα άρθρα)· Β: μόνο ο κωδικός (έλεγχος εγκυρότητας).
COA_REQUIRED = {"G": set(COA_COLUMNS), "B": {"Κωδικός"}}
COA_DIGITS = {"G": 10, "B": 6}
COA_FORMAT = {"G": "xx-xx-xx-xxxx", "B": "xx-xxxx"}
COA_LABEL = {"G": "Γ Κατηγορίας", "B": "Β Κατηγορίας"}

# (επικεφαλίδα, υποχρεωτικό, περιγραφή) — ίδια σειρά με τα πραγματικά αρχεία client_db
CLIENT_TEMPLATE_COLUMNS = (
    ("Κωδ. Συναλλασσόμενου", True, "Ακέραιος αριθμός, μοναδικός ανά συναλλασσόμενο (ο κωδικός που έχει στο λογιστικό πρόγραμμα)."),
    ("ΑΦΜ", True, "ΑΦΜ των 9 ψηφίων (κρατούνται τα αρχικά μηδενικά). Για εξωτερικό: ο ΑΦΜ/VAT με το πρόθεμα χώρας (π.χ. CY10245035I)."),
    ("Επωνυμία", True, "Επωνυμία ή ονοματεπώνυμο όπως στο λογιστικό πρόγραμμα."),
    ("Κωδ. ΔΟΥ", False, "Κωδικός ΔΟΥ (4 ψηφία), προαιρετικός."),
    ("Πόλη", True, "Πόλη."),
    ("Διεύθυνση", True, "Οδός και αριθμός."),
    ("Fax", False, "Προαιρετικό."),
    ("Τηλέφωνο", True, "Τηλέφωνο επικοινωνίας (μπορεί να μείνει κενό σε συναλλασσόμενο χωρίς τηλέφωνο)."),
    ("ΤΚ", True, "Ταχυδρομικός κώδικας."),
    ("Επάγγελμα", False, "Προαιρετικό."),
)

MAX_ROWS_LISTED = 15


# --------------------------------------------------------------------------- #
# Βασικά
# --------------------------------------------------------------------------- #
def norm_category(category: Any) -> str:
    s = str(category or "").strip().upper()
    return "G" if s.startswith("G") or s.startswith("Γ") else "B"


def safe_vat(vat: Any) -> str:
    return re.sub(r"[^0-9A-Za-z]", "", str(vat or "").strip())


def group_coa_path(base_dir: str, category: str) -> str:
    return os.path.join(base_dir, "chart_of_accounts_g.xlsx" if norm_category(category) == "G" else "chart_of_accounts_b.xlsx")


def credential_coa_path(base_dir: str, category: str, vat: str) -> Optional[str]:
    v = safe_vat(vat)
    if not v:
        return None
    return os.path.join(base_dir, COA_DIRNAME, f"{v}_{norm_category(category)}.xlsx")


def meta_path_for(path: str) -> str:
    return path + ".meta.json"


def resolve_coa_path(base_dir: str, category: str, vat: Any = None) -> Tuple[Optional[str], Optional[str]]:
    """(path, scope) — scope: 'credential' | 'group' | None όταν δεν υπάρχει σχέδιο."""
    cat = norm_category(category)
    cred = credential_coa_path(base_dir, cat, vat) if vat else None
    if cred and os.path.exists(cred):
        return cred, "credential"
    grp = group_coa_path(base_dir, cat)
    if os.path.exists(grp):
        return grp, "group"
    return None, None


def _read_meta(path: str) -> Dict[str, Any]:
    try:
        with open(meta_path_for(path), "r", encoding="utf-8") as fh:
            data = json.load(fh)
        return data if isinstance(data, dict) else {}
    except Exception:
        return {}


def scope_status(path: Optional[str], default_name: str = "") -> Dict[str, Any]:
    if not path or not os.path.exists(path):
        return {"exists": False}
    meta = _read_meta(path)
    return {
        "exists": True,
        "filename": meta.get("original_filename") or meta.get("filename") or default_name or os.path.basename(path),
        "uploaded_at": meta.get("uploaded_at", ""),
        "account_count": meta.get("account_count", 0),
    }


def coa_status(base_dir: str, category: str, vat: Any = None) -> Dict[str, Any]:
    """Κατάσταση και των δύο επιπέδων + ποιο είναι σε χρήση για το credential."""
    cat = norm_category(category)
    grp = scope_status(group_coa_path(base_dir, cat))
    cred_path = credential_coa_path(base_dir, cat, vat) if vat else None
    cred = scope_status(cred_path) if cred_path else {"exists": False}
    effective = "credential" if cred.get("exists") else ("group" if grp.get("exists") else None)
    return {"category": cat, "group": grp, "credential": cred, "effective": effective}


def remove_coa(base_dir: str, category: str, vat: Any = None) -> bool:
    """Αφαιρεί το ειδικό σχέδιο του credential (vat) ή, χωρίς vat, το κοινό της ομάδας."""
    path = credential_coa_path(base_dir, category, vat) if vat else group_coa_path(base_dir, category)
    removed = False
    for p in (path, meta_path_for(path) if path else None):
        if p and os.path.exists(p):
            os.remove(p)
            removed = True
    return removed


def store_credential_coa(base_dir: str, category: str, vat: str, data: bytes, meta_extra: Dict[str, Any],
                         keep_backups: int = 5) -> Dict[str, Any]:
    """Αποθηκεύει το ειδικό σχέδιο του credential (με αντίγραφο ασφαλείας του προηγούμενου)."""
    cat = norm_category(category)
    path = credential_coa_path(base_dir, cat, vat)
    if not path:
        raise ValueError("Λείπει ο ΑΦΜ του credential.")
    os.makedirs(os.path.dirname(path), exist_ok=True)
    ts = datetime.utcnow().strftime("%Y%m%dT%H%M%SZ")
    if os.path.exists(path):
        os.replace(path, f"{path}.bak.{ts}")
    with open(path, "wb") as fh:
        fh.write(data)
    # κράτα μόνο τα πιο πρόσφατα αντίγραφα αυτού του αρχείου
    folder, base = os.path.dirname(path), os.path.basename(path)
    backups = sorted(
        (n for n in os.listdir(folder) if n.startswith(base + ".bak.")),
        key=lambda n: os.path.getmtime(os.path.join(folder, n)), reverse=True,
    )
    for old in backups[keep_backups:]:
        try:
            os.remove(os.path.join(folder, old))
        except OSError:
            pass
    meta = {
        "filename": os.path.basename(path),
        "uploaded_at": datetime.utcnow().replace(microsecond=0).isoformat() + "Z",
        "category": cat,
        "scope": "credential",
        "vat": safe_vat(vat),
    }
    meta.update(meta_extra or {})
    with open(meta_path_for(path), "w", encoding="utf-8") as fh:
        json.dump(meta, fh, ensure_ascii=False, indent=2)
    return meta


# --------------------------------------------------------------------------- #
# Αναφορά ελέγχου
# --------------------------------------------------------------------------- #
class Report:
    """Συγκεντρώνει ευρήματα ελέγχου: errors (μπλοκάρουν), warnings (πρέπει να τα δει ο χρήστης), info."""

    def __init__(self) -> None:
        self.errors: List[Dict[str, Any]] = []
        self.warnings: List[Dict[str, Any]] = []
        self.info: List[Dict[str, Any]] = []

    def add(self, level: str, code: str, message: str, rows: Optional[Iterable[int]] = None,
            count: Optional[int] = None) -> None:
        rows_l = sorted(set(int(r) for r in (rows or [])))
        item: Dict[str, Any] = {"code": code, "message": message}
        if rows_l:
            item["rows"] = rows_l[:MAX_ROWS_LISTED]
            item["count"] = count if count is not None else len(rows_l)
            if len(rows_l) > MAX_ROWS_LISTED:
                item["rows_more"] = len(rows_l) - MAX_ROWS_LISTED
        elif count is not None:
            item["count"] = count
        getattr(self, {"error": "errors", "warning": "warnings", "info": "info"}[level]).append(item)

    def as_dict(self, **stats: Any) -> Dict[str, Any]:
        return {"ok": not self.errors, "errors": self.errors, "warnings": self.warnings, "info": self.info, "stats": stats}


def _fold(text: Any) -> str:
    s = unicodedata.normalize("NFD", str(text or "").strip().lower())
    return "".join(ch for ch in s if not unicodedata.combining(ch)).replace("ς", "σ")


def _excel_row(idx: int) -> int:
    return int(idx) + 2  # γραμμή 1 = επικεφαλίδες


def _clean_df(df):
    df = df.copy()
    df.columns = [str(c).strip() for c in df.columns]
    df = df.fillna("")
    return df


# --------------------------------------------------------------------------- #
# Έλεγχος λογιστικού σχεδίου
# --------------------------------------------------------------------------- #
def _digits(s: str) -> str:
    return "".join(ch for ch in s if ch.isdigit())


def _dashed(digits: str, category: str) -> str:
    if norm_category(category) == "G":
        return f"{digits[0:2]}-{digits[2:4]}-{digits[4:6]}-{digits[6:10]}"
    return f"{digits[0:2]}-{digits[2:6]}"


def validate_coa_df(df, category: str) -> Dict[str, Any]:
    cat = norm_category(category)
    rep = Report()
    df = _clean_df(df)
    cols = set(df.columns)
    missing = sorted(COA_REQUIRED[cat] - cols)
    if missing:
        rep.add("error", "missing_columns", "Λείπουν υποχρεωτικές στήλες: " + ", ".join(missing))
        return rep.as_dict(rows=len(df), detected_columns=sorted(cols))
    if len(df) == 0:
        rep.add("error", "empty", "Το αρχείο δεν περιέχει λογαριασμούς.")
        return rep.as_dict(rows=0)

    want = COA_DIGITS[cat]
    codes = df["Κωδικός"].astype(str).str.strip()
    empty_rows, full_rows, group_rows, no_dash_rows = [], [], [], []
    seen: Dict[str, List[int]] = {}
    code_set = set()
    for idx, code in codes.items():
        if not code:
            empty_rows.append(_excel_row(idx))
            continue
        code_set.add(code)
        seen.setdefault(code, []).append(_excel_row(idx))
        d = _digits(code)
        if len(d) == want and re.fullmatch(r"[\d\-\s]+", code) and ("-" not in code or code == _dashed(d, cat)):
            full_rows.append(_excel_row(idx))
            if code != _dashed(d, cat):
                no_dash_rows.append(_excel_row(idx))
        else:
            group_rows.append(_excel_row(idx))

    if not full_rows:
        rep.add("error", "no_valid_codes",
                f"Δεν βρέθηκε κανένας λογαριασμός της μορφής {COA_FORMAT[cat]} ({want} ψηφία) για {COA_LABEL[cat]}. "
                "Ελέγξτε ότι ανεβάσατε το σχέδιο της σωστής κατηγορίας.")
        return rep.as_dict(rows=len(df), accounts=0, group_accounts=len(group_rows))

    dup = [r for rows in seen.values() if len(rows) > 1 for r in rows]
    if dup:
        rep.add("warning", "duplicate_codes", "Υπάρχουν κωδικοί που εμφανίζονται περισσότερες από μία φορά "
                "(θα χρησιμοποιηθεί η πρώτη εμφάνιση).", dup)
    if no_dash_rows:
        rep.add("warning", "code_format",
                f"Κωδικοί χωρίς την τυπική μορφή {COA_FORMAT[cat]} (π.χ. χωρίς παύλες). Οι λογαριασμοί που συμπληρώνετε "
                "στις ρυθμίσεις έχουν παύλες, οπότε δεν θα ταυτιστούν με αυτές τις γραμμές.", no_dash_rows)

    vat_rate_col = "Ποσοστό ΦΠΑ" in cols
    vat_acct_col = "Λογαριασμός ΦΠΑ" in cols
    bad_rate, missing_vat_acct, bad_vat_fmt, rate_no_acct = [], [], [], []
    if vat_rate_col or vat_acct_col:
        for idx in df.index:
            row_no = _excel_row(idx)
            code = str(df.at[idx, "Κωδικός"]).strip()
            if not code:
                continue
            rate_raw = str(df.at[idx, "Ποσοστό ΦΠΑ"]).strip() if vat_rate_col else ""
            rate_val: Optional[float] = None
            if rate_raw:
                try:
                    rate_val = float(rate_raw.replace("%", "").replace(",", "."))
                    if not (0 <= rate_val <= 100):
                        raise ValueError
                except ValueError:
                    bad_rate.append(row_no)
                    rate_val = None
            vat_acct = str(df.at[idx, "Λογαριασμός ΦΠΑ"]).strip() if vat_acct_col else ""
            if vat_acct:
                if cat == "G" and (len(_digits(vat_acct)) != want or not re.fullmatch(r"[\d\-\s]+", vat_acct)):
                    bad_vat_fmt.append(row_no)
                elif vat_acct not in code_set:
                    missing_vat_acct.append(row_no)
            elif cat == "G" and rate_val and rate_val > 0 and row_no in set(full_rows):
                rate_no_acct.append(row_no)
    if bad_rate:
        rep.add("warning", "bad_vat_rate", "Μη έγκυρο «Ποσοστό ΦΠΑ» (αναμένεται αριθμός 0–100, π.χ. 24).", bad_rate)
    if bad_vat_fmt:
        rep.add("warning", "vat_account_format",
                f"«Λογαριασμός ΦΠΑ» με μη έγκυρη μορφή (αναμένεται {COA_FORMAT[cat]}).", bad_vat_fmt)
    if missing_vat_acct:
        rep.add("warning", "vat_account_missing",
                "«Λογαριασμός ΦΠΑ» που δεν υπάρχει στη στήλη «Κωδικός» του ίδιου αρχείου.", missing_vat_acct)
    if rate_no_acct:
        rep.add("info", "rate_without_vat_account",
                "Λογαριασμοί με Ποσοστό ΦΠΑ > 0 χωρίς «Λογαριασμό ΦΠΑ» — δεν θα προστεθεί αυτόματα γραμμή ΦΠΑ.",
                rate_no_acct)
    if empty_rows:
        rep.add("info", "empty_code", "Γραμμές χωρίς κωδικό (αγνοούνται).", empty_rows)
    if group_rows:
        rep.add("info", "group_accounts",
                f"{len(group_rows)} γραμμές είναι ομαδικοί/μη τελικοί λογαριασμοί (δεν έχουν {want} ψηφία) — "
                "δεν προτείνονται στα πεδία λογαριασμών.", count=len(group_rows))
    return rep.as_dict(rows=len(df), accounts=len(full_rows), group_accounts=len(group_rows))


# --------------------------------------------------------------------------- #
# Έλεγχος συναλλασσομένων (client_db)
# --------------------------------------------------------------------------- #
def afm_checksum_ok(afm: str) -> bool:
    """Έλεγχος ΑΦΜ (αλγόριθμος mod 11) για ΑΦΜ 9 ψηφίων."""
    if not re.fullmatch(r"\d{9}", afm or "") or afm == "000000000":
        return False
    total = sum(int(afm[i]) * (2 ** (8 - i)) for i in range(8))
    return (total % 11) % 10 == int(afm[8])


def _find_col(columns: List[str], needles: Iterable[str], exact: Iterable[str] = ()) -> Optional[str]:
    folded = {c: _fold(c) for c in columns}
    for c, f in folded.items():
        if f in set(exact):
            return c
    for c, f in folded.items():
        if any(n in f for n in needles):
            return c
    return None


def validate_clients_df(df) -> Dict[str, Any]:
    rep = Report()
    df = _clean_df(df)
    columns = list(df.columns)
    col_afm = _find_col(columns, ("αφμ", "afm", "vat"))
    col_id = _find_col(columns, ("συναλλ",), exact=("κωδικοσ", "κωδ", "κωδ.", "id", "κωδ. πελατη", "κωδικοσ πελατη"))
    col_name = _find_col(columns, ("επωνυ", "name"))
    if len(df) == 0:
        rep.add("error", "empty", "Το αρχείο δεν περιέχει γραμμές.")
        return rep.as_dict(rows=0)

    no_afm, bad_afm, foreign_afm, bad_id, no_name = [], [], [], [], []
    afm_rows: Dict[str, List[int]] = {}
    id_rows: Dict[str, List[int]] = {}
    for idx in df.index:
        row_no = _excel_row(idx)
        afm_raw = str(df.at[idx, col_afm]).strip() if col_afm else ""
        afm = re.sub(r"\s+", "", afm_raw).upper()
        if afm.startswith("EL") and re.fullmatch(r"EL\d{9}", afm):
            afm = afm[2:]
        if not afm:
            no_afm.append(row_no)
        else:
            afm_rows.setdefault(afm, []).append(row_no)
            if re.fullmatch(r"\d{9}", afm):
                if not afm_checksum_ok(afm):
                    bad_afm.append(row_no)
            elif re.fullmatch(r"\d{1,8}", afm):
                bad_afm.append(row_no)  # ελληνικός ΑΦΜ με χαμένα αρχικά μηδενικά ή ελλιπής
            else:
                foreign_afm.append(row_no)
        if col_id:
            raw_id = str(df.at[idx, col_id]).strip()
            try:
                cid = int(float(raw_id))
                id_rows.setdefault(str(cid), []).append(row_no)
            except (ValueError, TypeError):
                if afm:
                    bad_id.append(row_no)
        if col_name and not str(df.at[idx, col_name]).strip() and afm:
            no_name.append(row_no)

    if bad_id:
        rep.add("warning", "bad_customer_code",
                "«Κωδ. Συναλλασσόμενου» κενός ή μη ακέραιος — τέτοιοι συναλλασσόμενοι παραλείπονται από τις εξαγωγές.", bad_id)
    if bad_afm:
        rep.add("warning", "invalid_afm",
                "ΑΦΜ που δεν περνά τον έλεγχο εγκυρότητας (9 ψηφία, αρχικά μηδενικά, ψηφίο ελέγχου).", bad_afm)
    dup_afm = [r for rows in afm_rows.values() if len(rows) > 1 for r in rows]
    if dup_afm:
        rep.add("warning", "duplicate_afm", "ΑΦΜ που εμφανίζεται σε περισσότερες από μία γραμμές.", dup_afm)
    dup_id = [r for rows in id_rows.values() if len(rows) > 1 for r in rows]
    if dup_id:
        rep.add("warning", "duplicate_customer_code", "Ίδιος «Κωδ. Συναλλασσόμενου» σε περισσότερες γραμμές.", dup_id)
    if no_name:
        rep.add("warning", "missing_name", "Συναλλασσόμενος χωρίς επωνυμία.", no_name)
    if no_afm:
        rep.add("info", "no_afm", "Γραμμές χωρίς ΑΦΜ (αγνοούνται στις εξαγωγές).", no_afm)
    if foreign_afm:
        rep.add("info", "foreign_afm", "ΑΦΜ εξωτερικού (με γράμματα) — δεν ελέγχονται με τον ελληνικό αλγόριθμο.", foreign_afm)
    return rep.as_dict(rows=len(df), unique_afm=len(afm_rows))


# --------------------------------------------------------------------------- #
# Πρότυπα Excel
# --------------------------------------------------------------------------- #
def _styles():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    thin = Side(style="thin", color="BFBFBF")
    return {
        "req": PatternFill("solid", fgColor="FDE9D9"),
        "opt": PatternFill("solid", fgColor="EDEDED"),
        "head_font": Font(bold=True),
        "wrap": Alignment(wrap_text=True, vertical="top"),
        "border": Border(left=thin, right=thin, top=thin, bottom=thin),
        "title": Font(bold=True, size=13),
        "italic": Font(italic=True, color="808080"),
    }


def _write_instructions(ws, title: str, lines: List[Tuple[str, str]], widths=(26, 100)) -> None:
    st = _styles()
    ws["A1"] = title
    ws["A1"].font = st["title"]
    row = 3
    for left, right in lines:
        ws.cell(row=row, column=1, value=left).font = st["head_font"]
        c = ws.cell(row=row, column=2, value=right)
        c.alignment = st["wrap"]
        ws.cell(row=row, column=1).alignment = st["wrap"]
        row += 1
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w


def build_coa_template_bytes(category: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation

    cat = norm_category(category)
    st = _styles()
    wb = Workbook()
    ws = wb.active
    ws.title = "ΛΟΓΙΣΤΙΚΟ ΣΧΕΔΙΟ"
    if cat == "G":
        cols = [("Κωδικός", True), ("Περιγραφή", True), ("Ποσοστό ΦΠΑ", True), ("Λογαριασμός ΦΠΑ", True)]
    else:
        cols = [("Κωδικός", True), ("Περιγραφή", False), ("Ποσοστό ΦΠΑ", False), ("Λογαριασμός ΦΠΑ", False)]
    for i, (name, req) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=i, value=name)
        c.font = st["head_font"]
        c.fill = st["req"] if req else st["opt"]
        c.border = st["border"]
    for col_letter, w in zip("ABCD", (20, 52, 16, 22)):
        ws.column_dimensions[col_letter].width = w
    ws.freeze_panes = "A2"
    for r in range(2, 3002):  # κείμενο ώστε να μη χάνονται μηδενικά/παύλες
        for col in ("A", "D"):
            ws[f"{col}{r}"].number_format = "@"
    dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="100", allow_blank=True,
                        showErrorMessage=True, errorTitle="Ποσοστό ΦΠΑ", error="Δώστε αριθμό από 0 έως 100 (π.χ. 24).")
    ws.add_data_validation(dv)
    dv.add("C2:C3001")

    fmt = COA_FORMAT[cat]
    ws2 = wb.create_sheet("Οδηγίες")
    lines: List[Tuple[str, str]] = [
        ("Κατηγορία βιβλίων", f"Λογιστικό σχέδιο {COA_LABEL[cat]} — κωδικοί της μορφής {fmt} ({COA_DIGITS[cat]} ψηφία)."),
        ("Πώς το συμπληρώνετε", "Συμπληρώστε τις γραμμές στο φύλλο «ΛΟΓΙΣΤΙΚΟ ΣΧΕΔΙΟ» (μία γραμμή ανά λογαριασμό) και "
         "ανεβάστε το αρχείο στις Ρυθμίσεις. Το πρώτο φύλλο είναι αυτό που διαβάζεται· μην αλλάξετε τις επικεφαλίδες."),
        ("Κωδικός" + (" (υποχρεωτικό)"), f"Ο λογαριασμός, π.χ. {'60-00-00-0000' if cat == 'G' else '60-0000'}. "
         "Οι ομαδικοί λογαριασμοί (π.χ. 60, 60-00) επιτρέπονται αλλά δεν προτείνονται στα πεδία."),
        ("Περιγραφή" + (" (υποχρεωτικό)" if cat == "G" else " (προαιρετικό)"), "Ονομασία λογαριασμού."),
        ("Ποσοστό ΦΠΑ" + (" (υποχρεωτικό)" if cat == "G" else " (προαιρετικό)"),
         "Αριθμός 0–100 (π.χ. 24). Χρησιμοποιείται για να προτείνεται ο σωστός λογαριασμός ανά συντελεστή ΦΠΑ."),
        ("Λογαριασμός ΦΠΑ" + (" (υποχρεωτικό)" if cat == "G" else " (προαιρετικό)"),
         "Ο λογαριασμός ΦΠΑ που αντιστοιχεί στον λογαριασμό της γραμμής (π.χ. 54-00-98-0024). "
         "Πρέπει να υπάρχει και ως γραμμή στη στήλη «Κωδικός»." + (" Αν μείνει κενός δεν προστίθεται αυτόματα γραμμή ΦΠΑ." if cat == "G" else "")),
        ("Έλεγχος κατά το ανέβασμα", "Γίνεται έλεγχος: υπάρξη στηλών, διπλότυποι κωδικοί, μορφή κωδικών, εγκυρότητα ποσοστού ΦΠΑ και "
         "ύπαρξη του λογαριασμού ΦΠΑ στο ίδιο αρχείο. Τα ευρήματα εμφανίζονται αμέσως μετά το ανέβασμα."),
        ("Μορφοποίηση", "Οι στήλες κωδικών είναι μορφής «Κείμενο» ώστε να μη χαθούν μηδενικά ή παύλες. Πορτοκαλί επικεφαλίδα = "
         "υποχρεωτική στήλη, γκρι = προαιρετική."),
        ("Σχέδιο ανά credential", "Μπορείτε να ανεβάσετε ξεχωριστό σχέδιο για κάθε credential (ΑΦΜ). Αν δεν υπάρχει, "
         "χρησιμοποιείται το κοινό σχέδιο της ομάδας."),
    ]
    _write_instructions(ws2, "Οδηγίες συμπλήρωσης", lines)

    ws3 = wb.create_sheet("Παράδειγμα")
    ws3["A1"] = "Ενδεικτικές γραμμές — ΜΗΝ ανεβάσετε αυτό το φύλλο· αντιγράψτε ό,τι χρειάζεστε στο πρώτο φύλλο."
    ws3["A1"].font = st["italic"]
    for i, (name, _req) in enumerate(cols, start=1):
        c = ws3.cell(row=3, column=i, value=name)
        c.font = st["head_font"]
        c.border = st["border"]
    if cat == "G":
        sample = [
            ("20-00-00-0000", "ΕΜΠΟΡΕΥΜΑΤΑ", "", ""),
            ("54-00-98-0024", "ΦΠΑ ΕΙΣΡΟΩΝ 24%", "", ""),
            ("60-00-00-0000", "ΑΓΟΡΕΣ ΕΜΠΟΡΕΥΜΑΤΩΝ", "24", "54-00-98-0024"),
            ("64-00-00-0000", "ΓΕΝΙΚΑ ΕΞΟΔΑ", "24", "54-00-98-0024"),
            ("54-00-01-0024", "ΦΠΑ ΠΩΛΗΣΕΩΝ 24%", "", ""),
            ("70-00-00-0000", "ΠΩΛΗΣΕΙΣ ΕΜΠΟΡΕΥΜΑΤΩΝ", "24", "54-00-01-0024"),
        ]
    else:
        sample = [
            ("20-0000", "ΕΜΠΟΡΕΥΜΑΤΑ", "", ""),
            ("60-0000", "ΑΓΟΡΕΣ ΕΜΠΟΡΕΥΜΑΤΩΝ", "24", ""),
            ("64-0000", "ΓΕΝΙΚΑ ΕΞΟΔΑ", "24", ""),
            ("70-0000", "ΠΩΛΗΣΕΙΣ ΕΜΠΟΡΕΥΜΑΤΩΝ", "24", ""),
        ]
    for r, row in enumerate(sample, start=4):
        for c, v in enumerate(row, start=1):
            ws3.cell(row=r, column=c, value=v).number_format = "@"
    for col_letter, w in zip("ABCD", (20, 52, 16, 22)):
        ws3.column_dimensions[col_letter].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_clients_template_bytes() -> bytes:
    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation

    st = _styles()
    wb = Workbook()
    ws = wb.active
    ws.title = "ΣΥΝΑΛΛΑΣΣΟΜΕΝΟΙ"
    widths = {"Κωδ. Συναλλασσόμενου": 22, "ΑΦΜ": 14, "Επωνυμία": 48, "Κωδ. ΔΟΥ": 11, "Πόλη": 18, "Διεύθυνση": 34,
              "Fax": 14, "Τηλέφωνο": 16, "ΤΚ": 9, "Επάγγελμα": 36}
    text_cols = {"ΑΦΜ", "Κωδ. ΔΟΥ", "ΤΚ", "Τηλέφωνο", "Fax"}
    for i, (name, req, _d) in enumerate(CLIENT_TEMPLATE_COLUMNS, start=1):
        c = ws.cell(row=1, column=i, value=name)
        c.font = st["head_font"]
        c.fill = st["req"] if req else st["opt"]
        c.border = st["border"]
        letter = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[letter].width = widths.get(name, 18)
        if name in text_cols:
            for r in range(2, 5002):
                ws[f"{letter}{r}"].number_format = "@"  # κρατά τα αρχικά μηδενικά του ΑΦΜ/ΤΚ
    ws.freeze_panes = "A2"
    dv = DataValidation(type="whole", operator="greaterThan", formula1="0", allow_blank=True, showErrorMessage=True,
                        errorTitle="Κωδ. Συναλλασσόμενου", error="Δώστε ακέραιο αριθμό μεγαλύτερο του 0.")
    ws.add_data_validation(dv)
    dv.add("A2:A5001")

    ws2 = wb.create_sheet("Οδηγίες")
    lines: List[Tuple[str, str]] = [
        ("Πώς το συμπληρώνετε", "Μία γραμμή ανά συναλλασσόμενο στο φύλλο «ΣΥΝΑΛΛΑΣΣΟΜΕΝΟΙ». Μην αλλάξετε τις επικεφαλίδες. "
         "Πορτοκαλί = υποχρεωτική στήλη, γκρι = προαιρετική."),
        ("Ενημέρωση βάσης", "Το ανέβασμα συγχωνεύεται με την υπάρχουσα βάση με βάση τον ΑΦΜ: υπάρχων ΑΦΜ ενημερώνεται, νέος ΑΦΜ προστίθεται. "
         "Δεν σβήνονται συναλλασσόμενοι που δεν περιλαμβάνονται στο αρχείο."),
    ]
    for name, req, desc in CLIENT_TEMPLATE_COLUMNS:
        lines.append((name + (" (υποχρεωτικό)" if req else " (προαιρετικό)"), desc))
    lines.append(("Έλεγχος κατά το ανέβασμα", "Ελέγχεται: εγκυρότητα ΑΦΜ (ψηφίο ελέγχου), διπλότυποι ΑΦΜ και κωδικοί, κενός ή μη ακέραιος "
                  "«Κωδ. Συναλλασσόμενου», συναλλασσόμενος χωρίς επωνυμία. Τα ευρήματα εμφανίζονται αμέσως μετά το ανέβασμα."))
    _write_instructions(ws2, "Οδηγίες συμπλήρωσης — Συναλλασσόμενοι", lines)

    ws3 = wb.create_sheet("Παράδειγμα")
    ws3["A1"] = "Ενδεικτικές γραμμές (πλασματικά στοιχεία) — ΜΗΝ ανεβάσετε αυτό το φύλλο."
    ws3["A1"].font = st["italic"]
    for i, (name, _req, _d) in enumerate(CLIENT_TEMPLATE_COLUMNS, start=1):
        c = ws3.cell(row=3, column=i, value=name)
        c.font = st["head_font"]
        c.border = st["border"]
        ws3.column_dimensions[c.column_letter].width = widths.get(name, 18)
    sample = [
        (1001, "099999999", "ΔΕΙΓΜΑ ΕΜΠΟΡΙΚΗ ΙΚΕ", "1101", "ΑΘΗΝΑ", "ΟΔΟΣ ΔΟΚΙΜΗΣ 1", "", "2100000000", "10431", "ΕΜΠΟΡΙΟ"),
        (1002, "CY10245035I", "EXAMPLE TRADING LTD", "", "ΛΕΥΚΩΣΙΑ", "EXAMPLE STR 5", "", "", "1010", ""),
    ]
    for r, row in enumerate(sample, start=4):
        for c, v in enumerate(row, start=1):
            ws3.cell(row=r, column=c, value=v)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
